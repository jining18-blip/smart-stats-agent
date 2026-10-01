# -*- coding: utf-8 -*-
"""보고서 문장 · p값 표기 · CV 판정 · 엑셀 그래프 · PDF 표 읽기 검증.

앱의 해당 함수 구간만 잘라 실행한다(스트림릿 화면 없이).
"""
from pathlib import Path
import io
import re

import numpy as np
import pandas as pd
import pytest
import statsmodels.api as sm
from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1] / "app.py"
_SRC = APP.read_text(encoding="utf-8")


def _seg(a, b):
    return _SRC[_SRC.index(a):_SRC.index(b)]


@pytest.fixture(scope="module")
def ns():
    import streamlit as st
    n = {"np": np, "pd": pd, "re": re, "io": io, "st": st}
    for a, b in [("def fmt_p", "def round_half_up"),
                 ("def clean_columns", "def validate_anova_data"),
                 ("def cv_grade", "def find_numeric_like"),
                 ("def _v1_star", "def _v1_copy_box"),
                 ("def _josa", "def interpret_anova"),
                 ("_XL_NOT_VALUE = (", "def dl_table")]:
        exec(_seg(a, b), n)
    return n


# ---------------------------------------------------------------- p값 표기
def test_fmt_p_small_values_use_less_than(ns):
    f = ns["fmt_p"]
    assert f(0.00001) == "p<0.001"
    assert f(0.0) == "p<0.001"
    assert f(0.0234) == "p=0.0234"
    assert f(0.05, sp=True, digits=3) == "p = 0.050"
    assert f(np.nan) == "p=-" and f(None) == "p=-"


def test_no_p_equals_zero_in_sentences(ns):
    """문장 f-string 안에 p=0.0000 이 나올 수 있는 형식이 남아 있지 않아야 한다."""
    assert not re.search(r"p\s*=\s*\{[^}]*:\.4f\}", _SRC)


# ---------------------------------------------------------------- CV
def test_cv_grade_flags_too_low(ns):
    g = ns["cv_grade"]
    assert g(0.4) == "확인 필요(너무 낮음)"
    assert g(5) == "매우 우수" and g(15) == "양호" and g(25) == "다소 높음" and g(40) == "재검토 필요"
    assert g(np.nan) == "-"


def test_anova_sentence_warns_on_tiny_cv(ns):
    means = pd.DataFrame({"mean": [620.0, 600.0]}, index=["처리1", "대조구"])
    txt = ns["report_sentence_anova"]("처리구", "수량", 0.01, means, {"처리1": "a", "대조구": "b"},
                                      ci={"CV": 0.3})
    assert "반복별 원자료" in txt and "p=0.0100" in txt
    txt2 = ns["report_sentence_anova"]("처리구", "수량", 0.0001, means, {"처리1": "a", "대조구": "b"},
                                       ci={"CV": 8.0})
    assert "매우 우수" in txt2 and "p<0.001" in txt2


# ---------------------------------------------------------------- 보고서 문장
def test_corr_sentence_centers_on_target(ns):
    pairs = [("초장", "수량", 0.82, 0.0001), ("착과수", "수량", -0.55, 0.01), ("초장", "착과수", 0.1, 0.6)]
    txt = ns["report_sentence_corr"](pairs, (30, 30), "Pearson", target="수량")
    assert "n=30" in txt
    assert "초장(r=0.82***)" in txt and "정(+)의 상관" in txt
    assert "착과수(r=-0.55*)" in txt and "부(−)의 상관" in txt
    assert "인과관계" in txt


def test_corr_sentence_none_significant(ns):
    txt = ns["report_sentence_corr"]([("a", "b", 0.1, 0.5)], (10, 12))
    assert "n=10~12" in txt and "인정되지 않았다" in txt


def test_regression_sentence_and_footnote(ns):
    rng = np.random.default_rng(0)
    x = rng.normal(70, 5, 40)
    y = 8 * x + rng.normal(0, 10, 40)
    m = sm.OLS(pd.Series(y, name="수량"), sm.add_constant(pd.DataFrame({"초장": x}))).fit()
    txt = ns["report_sentence_reg"](m, "수량", ["초장"])
    assert "단순회귀분석" in txt and "n=40" in txt and "유의하였으며" in txt
    assert "초장이(가) 클수록 수량이(가) 증가" in txt
    foot = ns["reg_footnote"](m, ["초장"])
    assert foot.startswith("* p<0.05") and "n = 40" in foot and "수정 R²" not in foot


def test_ml_grade_and_sentence_units(ns):
    assert ns["ml_grade"](0.8, True)[0] == "좋음"
    assert ns["ml_grade"](0.6, True)[0] == "보통"
    assert ns["ml_grade"](0.2, True)[0] == "약함"
    assert ns["ml_grade"](-0.3, True)[0] == "매우 약함"
    assert ns["ml_grade"](0.9, False)[0] == "좋음"
    txt = ns["report_sentence_ml"]("수량(kg/10a)", "랜덤포레스트", True, 0.72, 40, 10,
                                   mae=12.3, rmse=15.1, top_vars=["초장", "착과수"])
    assert "±12.3 kg/10a" in txt and "'좋음'" in txt and "초장, 착과수" in txt
    txt2 = ns["report_sentence_ml"]("등급", "랜덤포레스트", False, 0.8, 40, 10, baseline=0.5)
    assert "80.0%" in txt2 and "50.0%" in txt2


def test_survey_sentences(ns):
    summ = pd.DataFrame({"문항": ["Q1", "Q2"], "평균": [4.2, 3.1], "긍정(%)": [80.0, 40.0]})
    txt = ns["report_sentence_likert"](summ, 0.85, 5, "긍정(%)", 60)
    assert "'Q1'이 4.20점" in txt and "0.850로 높은" in txt
    mc = pd.DataFrame({"문항": ["선호"] * 3, "응답": ["관수", "방제", "통합"],
                       "빈도": [30, 20, 10], "비율(%)": [50.0, 33.3, 16.7]})
    assert "'관수'이 50.0%(30명)" in ns["report_sentence_mc"](mc)
    mr = pd.DataFrame({"응답 항목": ["병해", "관수"], "응답 수": [40, 20], "응답률(%)": [66.7, 33.3]})
    assert "1인당 평균 1.0개" in ns["report_sentence_mr"]("관심분야", mr, 60)


def test_crosstab_sentence_uses_row_percent(ns):
    ct = pd.DataFrame([[8, 2], [1, 9]], index=["농가", "연구소"], columns=["예", "아니오"])
    txt = ns["report_sentence_crosstab"]("응답자", "재사용", ct, chi=(9.9, 1, 0.0005, 0.0))
    assert "농가(n=10)은 '예' 응답이 80.0%" in txt
    assert "연구소(n=10)은 '아니오' 응답이 90.0%" in txt
    assert "유의한 차이가 있었다" in txt and "p<0.001" in txt
    low = ns["report_sentence_crosstab"]("응답자", "재사용", ct, chi=(1.0, 1, 0.3, 50.0))
    assert "참고용" in low


# ---------------------------------------------------------------- 엑셀 그래프
def _charts(raw):
    wb = load_workbook(io.BytesIO(raw))
    return wb, {ws.title: [type(c).__name__ for c in ws._charts] for ws in wb.worksheets}


@pytest.mark.parametrize("kind,expect", [
    ("bar", "BarChart"), ("barh", "BarChart"), ("stacked", "BarChart"), ("stacked100", "BarChart"),
    ("line", "LineChart"), ("pie", "PieChart"), ("donut", "DoughnutChart"), ("scatter", "ScatterChart"),
])
def test_spec_chart_kinds(ns, kind, expect):
    data = pd.DataFrame({"항목": ["A", "B", "C"], "값1": [1.0, 2.0, 3.0], "값2": [2.0, 1.0, 4.0]})
    if kind == "scatter":
        data = pd.DataFrame({"초장": [60.0, 65, 70, 75], "수량": [500.0, 540, 590, 610]})
    raw = ns["make_xlsx"](data, "시험", chart_spec=ns["xl_chart"](kind, data, title="t"))
    _, ch = _charts(raw)
    assert expect in sum(ch.values(), [])


def test_heatmap_uses_color_scale_not_chart(ns):
    c = pd.DataFrame([[1, .5], [.5, 1]], index=["a", "b"], columns=["a", "b"]).reset_index()
    raw = ns["make_xlsx"](c, "상관", chart_spec=ns["xl_chart"]("heatmap", c, value_range=(-1, 1)))
    wb, ch = _charts(raw)
    assert not sum(ch.values(), [])
    assert any(len(ws.conditional_formatting) for ws in wb.worksheets)


def test_no_chart_subtitle_does_not_mention_graph(ns):
    d = pd.DataFrame({"항목": ["A", "B"], "값": [1.0, 2.0]})
    raw = ns["make_xlsx"](d, "표만", chart_spec=[])
    wb, ch = _charts(raw)
    assert not sum(ch.values(), [])
    texts = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows(max_row=3) for c in row
                     if c.value)
    assert "그래프" not in texts


def test_default_bar_chart_path_unchanged(ns):
    d = pd.DataFrame({"처리구": ["A", "B", "C"], "평균": [1.0, 2.0, 3.0]})
    _, ch = _charts(ns["make_xlsx"](d, "기본"))
    assert "BarChart" in sum(ch.values(), [])


def test_multi_sheet_keeps_spec_charts(ns):
    d = pd.DataFrame({"응답": ["예", "아니오"], "농가": [8, 2], "연구소": [1, 9]})
    blocks = [{"caption": "교차", "table": d, "xlsx_chart": [ns["xl_chart"]("stacked", d)]},
              {"caption": "표만", "table": d, "xlsx_chart": []}]
    _, ch = _charts(ns["make_xlsx_multi"](blocks, doc_title="t"))
    assert sum(len(v) for v in ch.values()) >= 1


# ---------------------------------------------------------------- PDF
def _pdf(bordered):
    rl = pytest.importorskip("reportlab")
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
    from reportlab.lib import colors
    rows = [["Trt", "Rep", "Yield"], ["A", "1", "600.5"], ["A", "2", "610.0"],
            ["B", "1", "650.2"], ["B", "2", "1,640.8"]]
    buf = io.BytesIO()
    t = Table(rows)
    if bordered:
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(buf, pagesize=A4).build([t])
    return buf.getvalue()


@pytest.mark.parametrize("bordered", [True, False])
def test_pdf_table_extraction(ns, bordered):
    pytest.importorskip("pdfplumber")
    tables, n_pages, n_text = ns["pdf_extract_tables"](_pdf(bordered))
    assert n_pages == 1 and n_text == 1 and tables
    d = tables[0]["df"]
    assert list(d.columns) == ["Trt", "Rep", "Yield"]
    assert len(d) == 4
    assert pd.api.types.is_numeric_dtype(d["Yield"]) and d["Yield"].iloc[-1] == pytest.approx(1640.8)


def test_pdf_page_to_png(ns):
    pytest.importorskip("pdfplumber")
    png = ns["pdf_page_png"](_pdf(True), 1, resolution=60)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
