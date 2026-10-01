# -*- coding: utf-8 -*-
"""데이터 점검(업로드 자료 확인) 검증.

앱의 점검 함수 구간만 잘라 실행한다. '무엇이 · 어디서 · 어떻게'가 맞게 나오는지,
그리고 정상 자료에 거짓 경고를 내지 않는지(처리1·1차 같은 이름) 확인한다.
"""
from pathlib import Path
import re

import numpy as np
import pandas as pd
import pytest

APP = Path(__file__).resolve().parents[1] / "app.py"
_SRC = APP.read_text(encoding="utf-8")


class _FakeSt:
    session_state = {"hdr_rows": 1}


@pytest.fixture(scope="module")
def ck():
    ns = {"np": np, "pd": pd, "re": re, "st": _FakeSt, "io": __import__("io")}
    for a, b in [("_BLOCK_KEYS =", "def _josa"),
                 ("def find_numeric_like", "def to_numeric_clean"),
                 ("_V1_TEMPLATES = {", "def _v1_data_readiness"),
                 ("_V1_SUMMARY_ROW_RE", "def _v1_render_checkup")]:
        exec(_SRC[_SRC.index(a):_SRC.index(b)], ns)
    return ns


def _levels(findings):
    return [(f["level"], f["title"]) for f in findings]


def _find(findings, word):
    hits = [f for f in findings if word in f["title"]]
    assert hits, _levels(findings)
    return hits[0]


def test_clean_rcbd_has_no_problems(ck):
    rng = np.random.default_rng(1)
    d = pd.DataFrame({"처리구": np.repeat(["대조구", "처리1", "처리2"], 3),
                      "반복": np.tile([1, 2, 3], 3),
                      "수량(kg/10a)": rng.normal(600, 20, 9).round(1)})
    out = ck["_v1_data_checkup"](d)
    assert not [f for f in out if f["level"] in ("error", "warn")], _levels(out)
    assert any("'처리구'" in f["title"] and f["level"] == "ok" for f in out)


def test_treatment_names_with_numbers_are_not_numeric(ck):
    d = pd.DataFrame({"처리구": ["대조구", "처리1", "처리2", "처리3"] * 2,
                      "조사시기": ["1차", "2차"] * 4, "수량": range(8)})
    assert ck["_v1_numeric_like"](d) == {}


def test_units_in_numbers_point_to_excel_rows(ck):
    d = pd.DataFrame({"처리구": ["A", "A", "B", "B"],
                      "수량(kg)": ["615", "620kg", "결측", "600"]})
    f = _find(ck["_v1_data_checkup"](d), "'수량(kg)' 열에 숫자가 아닌 값")
    assert f["level"] == "warn"
    assert "3행 '620kg'" in f["detail"] and "4행 '결측'" in f["detail"]      # 머리글이 1행이므로 +2


def test_title_row_instead_of_header(ck):
    d = pd.DataFrame([["처리구", "반복", "수량"], ["A", 1, 600], ["B", 1, 610]],
                     columns=["2025 고추 시험", "열", "열_2"])
    f = _find(ck["_v1_data_checkup"](d), "첫 행이 변수명이 아닌")
    assert f["level"] == "error"


def test_summary_row_and_label_clash(ck):
    d = pd.DataFrame({"처리구": ["대조구", "대조구 ", "처리1", "처리1", "평균"],
                      "수량": [600, 610, 650, 640, 625]})
    out = ck["_v1_data_checkup"](d)
    s = _find(out, "요약 행")
    assert s["level"] == "error" and "6행" in s["detail"]
    c = _find(out, "같은 이름이 다르게")
    assert "'대조구'" in c["detail"] and "'대조구 '" in c["detail"]


def test_missing_values_listed_by_column_and_row(ck):
    d = pd.DataFrame({"처리구": ["A", "A", "B", "B"], "수량": [600, np.nan, 610, np.nan]})
    f = _find(ck["_v1_data_checkup"](d), "빈칸")
    assert "'수량' 2칸" in f["detail"] and "3행" in f["detail"] and "5행" in f["detail"]


def test_leftover_template_examples(ck):
    cols = ck["_V1_TEMPLATES"]["일반 포장시험(처리×반복)"]
    ex = ck["_V1_TEMPLATE_EXAMPLES"]["일반 포장시험(처리×반복)"]
    d = pd.DataFrame(ex[:2] + [["처리2", 1, 80.0, 350.0, 700.0]], columns=cols)
    f = _find(ck["_v1_data_checkup"](d), "예시 행")
    assert f["level"] == "error" and "2행" in f["detail"] and "3행" in f["detail"]


def test_single_replicate_treatment(ck):
    d = pd.DataFrame({"처리구": ["A", "A", "A", "B", "B", "B", "C"],
                      "수량": [1, 2, 3, 4, 5, 6, 7.0]})
    f = _find(ck["_v1_data_checkup"](d), "반복이 1개뿐")
    assert "C" in f["detail"]


def test_wide_layout_detected(ck):
    d = pd.DataFrame({"반복": [1, 2, 3], "대조구": [500, 510, 505], "처리1": [560, 555, 570]})
    out = ck["_v1_data_checkup"](d)
    _find(out, "가로로 펼쳐진")


def test_id_column_not_used_as_measurement(ck):
    d = pd.DataFrame({"개체번호": range(1, 11), "초장": np.linspace(50, 70, 10)})
    assert ck["_v1_id_like_cols"](d) == ["개체번호"]


def test_template_workbook_has_examples_and_guide(ck):
    from openpyxl import load_workbook
    import io
    name = "일반 포장시험(처리×반복)"
    raw = ck["_v1_template_bytes"](ck["_V1_TEMPLATES"][name], ck["_V1_TEMPLATE_EXAMPLES"][name])
    wb = load_workbook(io.BytesIO(raw))
    assert wb.sheetnames == ["입력자료", "작성방법"]
    ws = wb["입력자료"]
    assert [c.value for c in ws[1]] == ck["_V1_TEMPLATES"][name]
    assert ws.cell(2, 1).value == "대조구" and ws.cell(2, 1).font.italic
    assert "작성방법" in ck["_V1_GUIDE_SHEETS"]
